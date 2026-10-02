from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "MANUAL_COMPLETO.md"
OUT = ROOT / "artifacts" / "Manual_Completo_PontoNorte.docx"
OUT.parent.mkdir(exist_ok=True)

NAVY = "123A55"
BLUE = "17678C"
GREEN = "17805B"
LIGHT = "EEF5F8"
TEXT = RGBColor(31, 47, 58)
MUTED = RGBColor(91, 111, 124)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.8)
sec.bottom_margin = Cm(1.7)
sec.left_margin = Cm(2.0)
sec.right_margin = Cm(2.0)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9.5)
styles["Normal"].font.color.rgb = TEXT
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Normal"].paragraph_format.line_spacing = 1.08
for name, size in [("Title", 28), ("Heading 1", 18), ("Heading 2", 13), ("Heading 3", 10.5)]:
    style = styles[name]
    style.font.name = "Aptos Display" if name != "Heading 3" else "Aptos"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
    style.paragraph_format.space_after = Pt(6)

if "Manual Lead" not in styles:
    lead = styles.add_style("Manual Lead", WD_STYLE_TYPE.PARAGRAPH)
    lead.font.name = "Aptos"
    lead.font.size = Pt(12)
    lead.font.color.rgb = MUTED
    lead.paragraph_format.space_after = Pt(12)

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margin(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("PontoNorte  •  Manual completo  •  ")
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = header.add_run("PONTO")
run.bold = True; run.font.size = Pt(9); run.font.color.rgb = RGBColor.from_string(NAVY)
run = header.add_run("NORTE")
run.bold = True; run.font.size = Pt(9); run.font.color.rgb = RGBColor.from_string(GREEN)
add_page_number(sec.footer.paragraphs[0])

cover = doc.add_paragraph()
cover.paragraph_format.space_before = Pt(80)
cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
mark = cover.add_run("◉")
mark.font.size = Pt(28); mark.font.color.rgb = RGBColor.from_string(GREEN)
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(5)
title_run = title.add_run("Manual completo do PontoNorte")
title_run.font.name = "Aptos Display"; title_run.font.size = Pt(28); title_run.bold = True; title_run.font.color.rgb = RGBColor(0,0,0)
sub = doc.add_paragraph(style="Manual Lead")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.add_run("Guia operacional para empresas, RH, líderes e colaboradores")
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.paragraph_format.space_before = Pt(18)
r = meta.add_run("Versão 1.3  •  2 de outubro de 2026\nPainel web, aplicativo Android e terminal corporativo")
r.font.size = Pt(10); r.font.color.rgb = MUTED
note = doc.add_table(rows=1, cols=1)
note.alignment = WD_ALIGN_PARAGRAPH.CENTER
note.autofit = False
note.columns[0].width = Cm(13)
cell = note.cell(0,0); cell.width = Cm(13); shade(cell, LIGHT); set_cell_margin(cell, 180, 220, 180, 220)
p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Controle de ponto multiempresa com jornadas, chamada, ocorrências, relatórios, horas extras, banco de horas e cobrança por funcionário ativo.")
r.font.size = Pt(10); r.font.color.rgb = RGBColor.from_string(NAVY)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

raw = SOURCE.read_text(encoding="utf-8").splitlines()
headings = [line[3:] for line in raw if line.startswith("## ")]
doc.add_heading("Sumário", level=1)
toc = doc.add_table(rows=0, cols=2)
toc.autofit = False
for idx, heading in enumerate(headings, 1):
    cells = toc.add_row().cells
    cells[0].width = Cm(1.2); cells[1].width = Cm(14.8)
    cells[0].text = f"{idx:02d}"
    cells[1].text = re.sub(r"^\d+\.\s*", "", heading)
    for c in cells:
        set_cell_margin(c, 45, 80, 45, 80)
        for p in c.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for rr in p.runs: rr.font.size = Pt(8.5); rr.font.color.rgb = MUTED
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def add_inline(paragraph, text):
    pattern = re.compile(r"(\*\*[^*]+\*\*|https?://\S+)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith("**"):
            rr = paragraph.add_run(token[2:-2]); rr.bold = True
        else:
            rr = paragraph.add_run(token.rstrip(".,;")); rr.font.color.rgb = RGBColor.from_string(BLUE); rr.underline = True
            if token[-1:] in ".,;": paragraph.add_run(token[-1])
        pos = m.end()
    if pos < len(text): paragraph.add_run(text[pos:])

def add_table(lines):
    rows = [[c.strip() for c in line.strip().strip("|").split("|")] for line in lines]
    if len(rows) > 1 and all(re.fullmatch(r":?-{3,}:?", c) for c in rows[1]):
        rows.pop(1)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = True
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i,j); cell.text = ""; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margin(cell)
            if i == 0: shade(cell, NAVY)
            p = cell.paragraphs[0]; add_inline(p, value); p.paragraph_format.space_after = Pt(0)
            for rr in p.runs:
                rr.font.size = Pt(8)
                if i == 0: rr.bold = True; rr.font.color.rgb = RGBColor(255,255,255)
        if i % 2 == 0 and i > 0:
            for cell in table.rows[i].cells: shade(cell, "F5F8FA")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

i = 0
skip_title = True
page_break_sections = {"4. Primeiro acesso da empresa", "17. Aba Relatórios", "23. Cobrança mensal por funcionário ativo"}
while i < len(raw):
    line = raw[i].rstrip()
    if skip_title and line.startswith("# "):
        skip_title = False; i += 1; continue
    if line.startswith("**Versão") or line.startswith("**Público") or line.startswith("**Plataforma") or line == "---":
        i += 1; continue
    if not line:
        i += 1; continue
    if line.startswith("|"):
        table_lines=[]
        while i < len(raw) and raw[i].startswith("|"):
            table_lines.append(raw[i]); i += 1
        add_table(table_lines); continue
    if line.startswith("## "):
        text=line[3:]
        if text in page_break_sections and len(doc.paragraphs) > 3:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading(text, level=1); i += 1; continue
    if line.startswith("### "):
        doc.add_heading(line[4:], level=2); i += 1; continue
    if line.startswith("#### "):
        doc.add_heading(line[5:], level=3); i += 1; continue
    if line.startswith("> "):
        table = doc.add_table(rows=1, cols=1); cell=table.cell(0,0); shade(cell, LIGHT); set_cell_margin(cell,130,160,130,160)
        p=cell.paragraphs[0]; add_inline(p,line[2:]); p.paragraph_format.space_after=Pt(0)
        i += 1; continue
    if re.match(r"^\d+\. ", line):
        match=re.match(r"^(\d+)\. (.*)$",line)
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.48); p.paragraph_format.first_line_indent=Cm(-.48)
        rr=p.add_run(match.group(1)+". "); rr.bold=True
        add_inline(p,match.group(2)); i += 1; continue
    if line.startswith("- "):
        p=doc.add_paragraph(style="List Bullet"); add_inline(p,line[2:]); i += 1; continue
    p=doc.add_paragraph(); add_inline(p,line); i += 1

for section in doc.sections:
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)

doc.core_properties.title = "Manual completo do PontoNorte"
doc.core_properties.subject = "Guia operacional da plataforma PontoNorte"
doc.core_properties.author = "PontoNorte"
doc.save(OUT)
print(OUT)
